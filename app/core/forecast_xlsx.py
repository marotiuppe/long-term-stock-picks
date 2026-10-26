"""The forecast as an Excel workbook, to check against the real prices afterwards.

Sheet "Forecast": one row per scanned stock and ETF with the last close, the low, middle and high case for the chosen period, and the chance of a
rise. The "Actual close" column is empty and yellow: type the real close in, and the columns beside it work out the real change, whether it fell
inside the low-to-high range, and whether the price rose. Sheet "Check" adds those up. Sheet "About" says what each column means.

The forecast is measured the way the scan measures it: the close on the scan date compared with the close `days` trading days later, so
"tomorrow" is the next trading day's close against today's.
"""
import datetime
import io

NAME = "Forecast"
FIRST = 2   # first data row (row 1 is the header)

COLUMNS = [  # (header, width)
    ("Symbol", 13), ("Name", 34), ("Exchange", 10), ("Type", 8), ("Call", 9),
    ("Last close date", 14), ("Last close (Rs)", 14), ("Forecast for the close of", 16),
    ("Low case % (10th)", 12), ("Middle % (median)", 12), ("High case % (90th)", 12),
    ("Low (Rs)", 11), ("Middle (Rs)", 11), ("High (Rs)", 11), ("Chance of a rise", 11),
    ("ACTUAL close (Rs): type it in", 17), ("Actual change %", 12), ("Inside low to high?", 13), ("Rose?", 8), ("Actual minus middle (points)", 14),
]
HEAD_FILL, INPUT_FILL = "0E2A25", "FFF2B3"


def _text(cell, value):
    """Store as text even when it starts with = + - @, so a company name can never be read as a formula."""
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"


def build(stocks, hz, days, target, generated, today=None):
    """Workbook bytes. stocks: the scan's stock dicts. hz: horizon key like "1d". days: trading days it spans. target: the date the forecast is for.
    generated: the scan's own timestamp text."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L

    today = today or datetime.date.today()
    wb = Workbook()
    ws = wb.active
    ws.title = NAME
    for c, (head, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=c, value=head)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=INPUT_FILL if c == 16 else HEAD_FILL)
        if c == 16:
            cell.font = Font(bold=True, color="000000")
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if c > 2 else "left")
        ws.column_dimensions[L(c)].width = width
    ws.row_dimensions[1].height = 46

    rows = sorted((s for s in stocks if isinstance(s.get("forecast"), dict) and s["forecast"].get(hz)), key=lambda s: -s.get("score", 0))
    for n, s in enumerate(rows):
        r = FIRST + n
        f, price = s["forecast"][hz], s["price"]
        asof = datetime.date.fromisoformat(s["asof"])
        _text(ws.cell(row=r, column=1), s["symbol"])
        _text(ws.cell(row=r, column=2), s.get("name") or s["symbol"])
        _text(ws.cell(row=r, column=3), s.get("exchange", "NSE"))
        _text(ws.cell(row=r, column=4), "ETF" if s.get("etf") else "Stock")
        _text(ws.cell(row=r, column=5), (s.get("verdict") or "").title())
        ws.cell(row=r, column=6, value=asof).number_format = "dd-mmm-yyyy"
        ws.cell(row=r, column=7, value=price).number_format = "#,##0.00"
        ws.cell(row=r, column=8, value=target).number_format = "dd-mmm-yyyy"
        for col, key in ((9, "p10"), (10, "p50"), (11, "p90")):
            ws.cell(row=r, column=col, value=round(f[key], 4)).number_format = "0.00%"
            ws.cell(row=r, column=col + 3, value=round(price * (1 + f[key]), 2)).number_format = "#,##0.00"
        ws.cell(row=r, column=15, value=round(f["pgain"], 3)).number_format = "0%"
        inp = ws.cell(row=r, column=16)
        inp.fill = PatternFill("solid", fgColor=INPUT_FILL)
        inp.number_format = "#,##0.00"
        ws.cell(row=r, column=17, value=f'=IF(P{r}="","",P{r}/G{r}-1)').number_format = "0.00%"
        ws.cell(row=r, column=18, value=f'=IF(P{r}="","",IF(AND(P{r}>=L{r},P{r}<=N{r}),"Yes","No"))')
        ws.cell(row=r, column=19, value=f'=IF(P{r}="","",IF(P{r}>G{r},"Yes","No"))')
        ws.cell(row=r, column=20, value=f'=IF(P{r}="","",Q{r}-J{r})').number_format = "0.00%"
        for col in (3, 4, 5, 6, 8, 15, 18, 19):
            ws.cell(row=r, column=col).alignment = Alignment(horizontal="center")
    last = FIRST + max(len(rows), 1) - 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{L(len(COLUMNS))}{last}"
    good, bad = PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE"), PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE")
    ws.conditional_formatting.add(f"R{FIRST}:S{last}", CellIsRule(operator="equal", formula=['"Yes"'], fill=good))
    ws.conditional_formatting.add(f"R{FIRST}:R{last}", CellIsRule(operator="equal", formula=['"No"'], fill=bad))

    # ---- Check: the totals, as formulas, so they update as the actual closes are typed in
    ck = wb.create_sheet("Check")
    rng = lambda col: f"{NAME}!{col}{FIRST}:{col}{last}"
    ck.column_dimensions["A"].width = 52
    ck.column_dimensions["B"].width = 16
    ck.column_dimensions["C"].width = 70
    ck["A1"] = "How the forecast did"
    ck["A1"].font = Font(bold=True, size=14)
    table = [
        ("Forecast period", f"{days} trading day{'s' if days != 1 else ''}", f"The close on {rows[0]['asof'] if rows else ''} against the close on {target.isoformat()}."),
        ("Stocks and ETFs in the sheet", f"=COUNTA({rng('A')})", ""),
        ("Actual closes you have typed in", f"=COUNT({rng('P')})", "Fill in the yellow column on the Forecast sheet."),
        ("Closed inside the low-to-high range", f'=IF(B5=0,"",COUNTIF({rng("R")},"Yes")/B5)', "The range is the 10th to 90th percentile, so about 80% should fall inside it. Far from 80% means the range is too narrow or too wide."),
        ("Rose (closed above the last close)", f'=IF(B5=0,"",COUNTIF({rng("S")},"Yes")/B5)', "Compare with the next line."),
        ("Average forecast chance of a rise", f'=IF(B5=0,"",SUMPRODUCT(({rng("P")}<>"")*{rng("O")})/B5)', "If the model is calibrated, the line above is close to this one."),
        ("Average actual change", f'=IFERROR(AVERAGE({rng("Q")}),"")', ""),
        ("Average forecast middle case", f'=IF(B5=0,"",SUMPRODUCT(({rng("P")}<>"")*{rng("J")})/B5)', ""),
        ("Buy calls with an actual close", f'=SUMPRODUCT(({rng("E")}="Buy")*({rng("P")}<>""))', "Only rows whose Call is Buy."),
        ("Buy calls that rose", f'=IF(B11=0,"",SUMPRODUCT(({rng("E")}="Buy")*({rng("S")}="Yes"))/B11)', "Compare with all stocks above: that gap is the whole point of a Buy label."),
    ]
    for i, (a, b, c) in enumerate(table, 3):
        ck.cell(row=i, column=1, value=a)
        ck.cell(row=i, column=2, value=b).alignment = Alignment(horizontal="right")
        ck.cell(row=i, column=3, value=c).alignment = Alignment(wrap_text=True, vertical="top")
    for r in (6, 7, 8, 12):
        ck.cell(row=r, column=2).number_format = "0.0%"
    for r in (9, 10):
        ck.cell(row=r, column=2).number_format = "0.00%"

    # ---- About
    ab = wb.create_sheet("About")
    ab.column_dimensions["A"].width = 120
    lines = [
        ("Forecast to check against the real prices", True),
        (f"Made from the scan of {generated}. Prices are end-of-day closes of the last trading day before the forecast.", False),
        (f"Period: {days} trading day{'s' if days != 1 else ''} ({hz}). Each row compares the close on its 'Last close date' with the close on {target.isoformat()}.", False),
        ("", False),
        ("How to check it", True),
        (f"1. After the market closes on {target.strftime('%d %b %Y')}, type each stock's real closing price into the yellow ACTUAL close column (NSE or your broker app).", False),
        ("2. Actual change %, Inside low to high?, Rose? and Actual minus middle fill themselves in. The Check sheet adds them up.", False),
        ("3. Use the filter buttons on the header row to look at only Buy calls, only ETFs and so on.", False),
        ("", False),
        ("What the columns mean", True),
        ("Low, Middle and High case: the 10th, 50th (median) and 90th percentile of what stocks with a similar score did over this period, scaled to the stock's own volatility. About 8 times in 10 the real change should land between Low and High.", False),
        ("Chance of a rise: the share of similar past cases that ended higher. It is a historical base rate, not a prediction.", False),
        ("Call: the scan's label (Buy, Watch, Wait, Avoid). A category of a statistical model, not an instruction to buy anything.", False),
        ("", False),
        ("Please read", True),
        ("For education only. Not investment advice and not a recommendation to buy or sell any stock or ETF. We are not responsible for decisions or losses. The model has no proven edge.", False),
        ("A forecast can be wrong by far more than the High or Low case. A few days of checking proves very little: use it to learn how wide the ranges really are.", False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = ab.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            c.font = Font(bold=True, size=13)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
